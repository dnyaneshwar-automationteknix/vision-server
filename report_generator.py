"""
Report Generator - PDF and Excel Export
Generate inspection reports in multiple formats
"""

import csv
from io import BytesIO, StringIO
from datetime import datetime, timedelta
from typing import Dict, List
from pathlib import Path


class ReportGenerator:
    """
    Generate reports in various formats
    PDF, Excel, CSV
    """
    
    def __init__(self, database):
        """
        Initialize report generator
        Args:
            database: QCDatabase instance
        """
        self.db = database
    
    # ==================== CSV EXPORT ====================
    
    def generate_csv_report(self, start_date: str = None, end_date: str = None) -> str:
        """
        Generate CSV report
        Returns: CSV string
        """
        # Get inspections
        inspections = self.db.get_inspections(
            limit=10000,
            start_date=start_date,
            end_date=end_date
        )
        
        # Create CSV
        output = StringIO()
        writer = csv.writer(output)
        
        # Headers
        writer.writerow([
            'Inspection ID',
            'Date/Time',
            'Product Code',
            'Product Name',
            'QR Data',
            'OCR Result',
            'Color Detected',
            'Status',
            'Blur Score',
            'Brightness Score',
            'Processing Time (s)',
            'Station ID',
            'Operator ID'
        ])
        
        # Data rows
        for insp in inspections:
            writer.writerow([
                insp.get('inspection_id', ''),
                insp.get('timestamp', ''),
                insp.get('product_code', ''),
                insp.get('product_name', ''),
                insp.get('qr_data', ''),
                insp.get('ocr_result', ''),
                insp.get('color_detected', ''),
                insp.get('status', ''),
                insp.get('blur_score', ''),
                insp.get('brightness_score', ''),
                insp.get('processing_time', ''),
                insp.get('station_id', ''),
                insp.get('operator_id', '')
            ])
        
        return output.getvalue()
    
    # ==================== EXCEL EXPORT ====================
    
    def generate_excel_report(self, start_date: str = None, end_date: str = None) -> bytes:
        """
        Generate Excel report
        Returns: Excel file bytes
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            
            # Create workbook
            wb = openpyxl.Workbook()
            
            # Summary Sheet
            ws_summary = wb.active
            ws_summary.title = "Summary"
            
            # Get statistics
            stats = self.db.get_statistics(start_date, end_date)
            
            # Title
            ws_summary['A1'] = 'QC Vision System - Inspection Report'
            ws_summary['A1'].font = Font(size=16, bold=True)
            ws_summary.merge_cells('A1:D1')
            
            # Date range
            ws_summary['A3'] = f'Report Generated: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'
            if start_date:
                ws_summary['A4'] = f'Date Range: {start_date} to {end_date or "Present"}'
            
            # Statistics
            ws_summary['A6'] = 'Statistics'
            ws_summary['A6'].font = Font(bold=True, size=14)
            
            stats_data = [
                ['Total Inspections', stats['total_inspections']],
                ['OK Count', stats['ok_count']],
                ['NOK Count', stats['nok_count']],
                ['Pass Rate', f"{stats['pass_rate']:.2f}%"]
            ]
            
            for i, (label, value) in enumerate(stats_data, start=7):
                ws_summary[f'A{i}'] = label
                ws_summary[f'B{i}'] = value
                ws_summary[f'A{i}'].font = Font(bold=True)
            
            # Inspections Sheet
            ws_inspections = wb.create_sheet("Inspections")
            
            # Headers
            headers = [
                'ID', 'Timestamp', 'Product Code', 'Product Name',
                'QR Data', 'OCR Result', 'Status', 
                'Blur Score', 'Brightness', 'Processing Time'
            ]
            
            for col, header in enumerate(headers, start=1):
                cell = ws_inspections.cell(row=1, column=col)
                cell.value = header
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            
            # Get inspections
            inspections = self.db.get_inspections(
                limit=10000,
                start_date=start_date,
                end_date=end_date
            )
            
            # Data rows
            for row_idx, insp in enumerate(inspections, start=2):
                ws_inspections.cell(row=row_idx, column=1, value=insp.get('inspection_id'))
                ws_inspections.cell(row=row_idx, column=2, value=insp.get('timestamp'))
                ws_inspections.cell(row=row_idx, column=3, value=insp.get('product_code'))
                ws_inspections.cell(row=row_idx, column=4, value=insp.get('product_name'))
                ws_inspections.cell(row=row_idx, column=5, value=insp.get('qr_data'))
                ws_inspections.cell(row=row_idx, column=6, value=insp.get('ocr_result'))
                ws_inspections.cell(row=row_idx, column=7, value=insp.get('status'))
                ws_inspections.cell(row=row_idx, column=8, value=insp.get('blur_score'))
                ws_inspections.cell(row=row_idx, column=9, value=insp.get('brightness_score'))
                ws_inspections.cell(row=row_idx, column=10, value=insp.get('processing_time'))
                
                # Color code status
                status_cell = ws_inspections.cell(row=row_idx, column=7)
                if insp.get('status') == 'OK':
                    status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                elif insp.get('status') == 'NOK':
                    status_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            
            # Adjust column widths
            for col in range(1, len(headers) + 1):
                ws_inspections.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
            
            # Save to bytes
            excel_file = BytesIO()
            wb.save(excel_file)
            excel_file.seek(0)
            
            return excel_file.getvalue()
        
        except ImportError:
            # If openpyxl not available, return CSV instead
            print("⚠️ openpyxl not installed. Install with: pip install openpyxl")
            return self.generate_csv_report(start_date, end_date).encode('utf-8')
    
    # ==================== PDF EXPORT ====================
    
    def generate_pdf_report(self, start_date: str = None, end_date: str = None) -> bytes:
        """
        Generate PDF report
        Returns: PDF file bytes
        """
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import letter, A4
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            
            # Create PDF buffer
            buffer = BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=A4)
            elements = []
            
            # Styles
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                textColor=colors.HexColor('#667eea'),
                spaceAfter=30,
                alignment=1  # Center
            )
            
            # Title
            title = Paragraph("QC Vision System<br/>Inspection Report", title_style)
            elements.append(title)
            
            # Date info
            date_text = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            if start_date:
                date_text += f"<br/>Period: {start_date} to {end_date or 'Present'}"
            
            date_para = Paragraph(date_text, styles['Normal'])
            elements.append(date_para)
            elements.append(Spacer(1, 20))
            
            # Statistics
            stats = self.db.get_statistics(start_date, end_date)
            
            stats_title = Paragraph("<b>Summary Statistics</b>", styles['Heading2'])
            elements.append(stats_title)
            elements.append(Spacer(1, 10))
            
            stats_data = [
                ['Metric', 'Value'],
                ['Total Inspections', str(stats['total_inspections'])],
                ['OK Count', str(stats['ok_count'])],
                ['NOK Count', str(stats['nok_count'])],
                ['Pass Rate', f"{stats['pass_rate']:.2f}%"]
            ]
            
            stats_table = Table(stats_data, colWidths=[3*inch, 2*inch])
            stats_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#667eea')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            elements.append(stats_table)
            elements.append(Spacer(1, 30))
            
            # Inspections table
            inspections_title = Paragraph("<b>Recent Inspections</b>", styles['Heading2'])
            elements.append(inspections_title)
            elements.append(Spacer(1, 10))
            
            # Get inspections (limit to 50 for PDF)
            inspections = self.db.get_inspections(limit=50, start_date=start_date, end_date=end_date)
            
            if inspections:
                insp_data = [['ID', 'Product', 'Status', 'Time']]
                
                for insp in inspections:
                    insp_data.append([
                        str(insp.get('inspection_id', 'N/A')),
                        insp.get('product_code', 'N/A')[:15],
                        insp.get('status', 'N/A'),
                        insp.get('timestamp', 'N/A')[:16]
                    ])
                
                insp_table = Table(insp_data, colWidths=[0.8*inch, 2*inch, 1*inch, 2*inch])
                insp_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#667eea')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.white),
                    ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                    ('FONTSIZE', (0, 1), (-1, -1), 8)
                ]))
                
                elements.append(insp_table)
            
            # Build PDF
            doc.build(elements)
            buffer.seek(0)
            
            return buffer.getvalue()
        
        except ImportError:
            print("⚠️ reportlab not installed. Install with: pip install reportlab")
            # Return simple text report
            return self._generate_text_report(start_date, end_date).encode('utf-8')
    
    def _generate_text_report(self, start_date: str = None, end_date: str = None) -> str:
        """Generate simple text report (fallback)"""
        stats = self.db.get_statistics(start_date, end_date)
        inspections = self.db.get_inspections(limit=20, start_date=start_date, end_date=end_date)
        
        report = f"""
QC VISION SYSTEM - INSPECTION REPORT
{'=' * 60}

Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
{'Date Range: ' + start_date + ' to ' + (end_date or 'Present') if start_date else ''}

SUMMARY STATISTICS
{'-' * 60}
Total Inspections:  {stats['total_inspections']}
OK Count:           {stats['ok_count']}
NOK Count:          {stats['nok_count']}
Pass Rate:          {stats['pass_rate']:.2f}%

RECENT INSPECTIONS
{'-' * 60}
"""
        
        for insp in inspections:
            report += f"\n#{insp.get('inspection_id')} | {insp.get('product_code', 'N/A'):<15} | {insp.get('status', 'N/A'):<6} | {insp.get('timestamp', 'N/A')}"
        
        report += f"\n\n{'=' * 60}\n"
        
        return report
    
    # ==================== UTILITY ====================
    
    def save_report(self, report_data: bytes, filename: str, output_dir: str = "reports"):
        """Save report to file"""
        output_path = Path(output_dir)
        output_path.mkdir(exist_ok=True)
        
        file_path = output_path / filename
        
        with open(file_path, 'wb') as f:
            f.write(report_data)
        
        return str(file_path)


# ==================== TEST/DEMO ====================

if __name__ == "__main__":
    from database import QCDatabase
    
    print("=" * 60)
    print("Report Generator - Test")
    print("=" * 60)
    
    # Setup test database
    db = QCDatabase("test_reports.db")
    
    # Add test data
    product_id = db.add_product("TEST-001", "Test Product")
    
    for i in range(10):
        db.add_inspection(
            product_id=product_id,
            qr_data=f"QR-{i:04d}",
            status='OK' if i % 3 != 0 else 'NOK',
            blur_score=90 + i,
            brightness_score=120 + i
        )
    
    print("\n✓ Test data created")
    
    # Create report generator
    generator = ReportGenerator(db)
    
    # Test CSV
    print("\n📄 Generating CSV report...")
    csv_report = generator.generate_csv_report()
    csv_path = generator.save_report(csv_report.encode('utf-8'), 'test_report.csv')
    print(f"  ✓ Saved: {csv_path}")
    
    # Test Excel
    print("\n📊 Generating Excel report...")
    try:
        excel_report = generator.generate_excel_report()
        excel_path = generator.save_report(excel_report, 'test_report.xlsx')
        print(f"  ✓ Saved: {excel_path}")
    except Exception as e:
        print(f"  ⚠️ Excel generation failed: {e}")
    
    # Test PDF
    print("\n📕 Generating PDF report...")
    try:
        pdf_report = generator.generate_pdf_report()
        pdf_path = generator.save_report(pdf_report, 'test_report.pdf')
        print(f"  ✓ Saved: {pdf_path}")
    except Exception as e:
        print(f"  ⚠️ PDF generation failed: {e}")
    
    db.close()
    
    print("\n✓ Report generation test complete!")
    print("\nTo enable all formats:")
    print("  pip install openpyxl reportlab")